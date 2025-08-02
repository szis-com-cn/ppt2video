import os
import json
import subprocess
from pptx import Presentation
from glob import glob

class PPT2JSON:
    """
    解析PPT文字和备注，并导出真实渲染的图片（unoconv+pdftoppm）
    """
    @staticmethod
    def extract_text_notes_and_images(pptx_path, output_dir="ppt_pages"):
        os.makedirs(output_dir, exist_ok=True)

        # 1. 转换为PDF并拆分为PNG
        PPT2JSON.convert_ppt_to_images(pptx_path, output_dir)
        
        # 获取所有PNG文件
        png_files = sorted(glob(os.path.join(output_dir, "output_slide-*.png")))

        # 2. 解析文本和备注
        prs = Presentation(pptx_path)
        slides_data = []
        for idx, (slide, img_path) in enumerate(zip(prs.slides, png_files), start=1):
            # 提取正文
            slide_text = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text.append(shape.text.strip())

            # 提取备注
            notes_text = ""
            if slide.has_notes_slide and slide.notes_slide:
                for shape in slide.notes_slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        notes_text += shape.text.strip() + "\n"
                notes_text = notes_text.strip()

            slides_data.append({
                "slide_number": idx,
                "slide_text": "\n".join(slide_text),
                "notes_text": notes_text,
                "image_path": img_path
            })

        print(f"[OK] PPT解析完成，共 {len(slides_data)} 页，并导出真实图片到 {output_dir}/")
        return slides_data

    @staticmethod
    def convert_ppt_to_images(pptx_path, output_dir):
        """
        使用 unoconv 将 PPTX 转为 PDF，再用 pdftoppm 拆分为 PNG
        """
        base_name = os.path.splitext(os.path.basename(pptx_path))[0]
        temp_pdf = os.path.join(output_dir, f"{base_name}.pdf")

        # Step 1: PPTX → PDF
        subprocess.run(["unoconv", "-f", "pdf", "-d", "presentation", "-o", temp_pdf, pptx_path], check=True)

        # Step 2: PDF → PNG
        subprocess.run(["pdftoppm", "-png", temp_pdf, os.path.join(output_dir, "output_slide")], check=True)

        # Step 3: 删除临时 PDF
        if os.path.exists(temp_pdf):
            os.remove(temp_pdf)

        print(f"[OK] 已导出所有幻灯片为 PNG，并删除临时 PDF")
